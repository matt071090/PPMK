"""Template master data iReapBiz (Petty Cash Category Asg, Warehouse Assignment, Employee Master Data).

Struktur field & dropdown diambil dari DB SAP B1 HANA 192.168.40.172 (default schema BFU_20260910).
Login memakai user/password datasource di CVPPMK/application.properties (DBADMIN).

    python generate_templates.py [SCHEMA]      mis. python generate_templates.py FFI_20260929

Konvensi STEM: baris 1 = label (* = wajib), baris 2 = nama field DTW, baris 3+ = contoh.
Petunjuk tiap kolom di sheet Keterangan, daftar pilihan dropdown di sheet Referensi.
"""
import sys
from datetime import date
from pathlib import Path

from hdbcli import dbapi
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HOST, PORT = "192.168.40.172", 30015
SCHEMA = sys.argv[1] if len(sys.argv) > 1 else "BFU_20260910"
HERE = Path(__file__).parent
PROPS = HERE.parent / "CVPPMK" / "application.properties"

LABEL_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FONT = Font(bold=True, color="FFFFFF")
FIELD_FILL = PatternFill("solid", fgColor="D9E1F2")
FIELD_FONT = Font(bold=True, color="1F4E78", italic=True)
NOTE_FONT = Font(italic=True, color="595959")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ROWS = 1000  # baris yang diberi dropdown


def props():
    out = {}
    for line in PROPS.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            k, v = line.split("=", 1)
            out[k.strip()] = v.strip()
    return out


def fetch_refs():
    p = props()
    con = dbapi.connect(address=HOST, port=PORT, user=p["spring.datasource.username"],
                        password=p["spring.datasource.password"], currentschema=SCHEMA)
    cur = con.cursor()

    def q(sql):
        cur.execute(sql)
        return [tuple("" if v is None else v for v in r) for r in cur.fetchall()]

    refs = {
        "company": q('SELECT "CompnyName" FROM OADM')[0][0],
        "accounts": q('SELECT "AcctCode", "AcctName" FROM OACT WHERE "Postable" = \'Y\' ORDER BY "AcctCode"'),
        "warehouses": q('SELECT "WhsCode", "WhsName" FROM OWHS WHERE "Inactive" = \'N\' ORDER BY "WhsCode"'),
        "authweb": q('SELECT "DocEntry", "DocNum", "U_Remarks" FROM "@STEM_AUTHWEBH" ORDER BY "DocEntry"'),
        "whsasg": q('SELECT "DocEntry", "DocNum", "U_Remarks" FROM "@STEM_WHSASGNWEBPH" ORDER BY "DocEntry"'),
        "pccatasg": q('SELECT "DocEntry", "DocNum", "U_Remarks" FROM "@STEM_PCCATASGH" ORDER BY "DocEntry"'),
        "pccat_lines": q('SELECT H."DocNum", D."LineId", D."U_AcctCode", D."U_PettyCashCatName" FROM "@STEM_PCCATASGD" D '
                         'JOIN "@STEM_PCCATASGH" H ON H."DocEntry" = D."DocEntry" ORDER BY H."DocNum", D."LineId"'),
        "whs_lines": q('SELECT H."DocNum", H."U_Remarks", D."LineId", D."U_WhsCode", D."U_WhsName" FROM "@STEM_WHSASGNWEBPD" D '
                       'JOIN "@STEM_WHSASGNWEBPH" H ON H."DocEntry" = D."DocEntry" ORDER BY H."DocNum", D."LineId"'),
    }
    for d in range(1, 6):
        refs[f"dim{d}"] = q(f'SELECT "OcrCode", "OcrName" FROM OOCR WHERE "DimCode" = {d} AND "Active" = \'Y\' ORDER BY "OcrCode"')
    con.close()
    return refs


def F(label, field, note, req=False, ref=None, enum=None, width=None):
    return dict(label=label, field=field, note=note, req=req, ref=ref, enum=enum, width=width)


def build(path, title, purpose, sheets, refs, ref_keys, extra_notes=()):
    """sheets: [(name, fields, sample_rows)], ref_keys: {key: (judul, kolom)} yang ditulis ke sheet Referensi."""
    wb = Workbook()
    wb.remove(wb.active)

    # sheet Referensi (dibuat dulu supaya range dropdown diketahui)
    ref_range, ref_blocks, col = {}, [], 1
    for key, (head, cols) in ref_keys.items():
        rows = refs.get(key, []) if not isinstance(key, tuple) else key[1]
        ref_blocks.append((col, head, cols, rows))
        n = max(len(rows), 1)
        ref_range[key if not isinstance(key, tuple) else key[0]] = f"'Referensi'!${get_column_letter(col)}$3:${get_column_letter(col)}${n + 2}"
        col += len(cols) + 1

    for name, fields, samples in sheets:
        ws = wb.create_sheet(name)
        for c, f in enumerate(fields, 1):
            for r, v, fill, font in ((1, f["label"] + (" *" if f["req"] else ""), LABEL_FILL, LABEL_FONT),
                                     (2, f["field"], FIELD_FILL, FIELD_FONT)):
                cell = ws.cell(r, c, v)
                cell.fill, cell.font, cell.border = fill, font, BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            letter = get_column_letter(c)
            ws.column_dimensions[letter].width = f["width"] or min(max(len(f["label"]), len(f["field"])) + 4, 45)
            for r in range(3, ROWS + 3):
                ws.cell(r, c).number_format = "@"  # semua kolom teks: kode seperti 001 / 6101001 tidak berubah
            src = ref_range.get(f["ref"]) if f["ref"] else None
            if f["enum"]:
                src = '"' + ",".join(f["enum"]) + '"'
            if src:
                # kode dokumen UDO yang baru dibuat lewat template ini belum ada di daftar: boleh diketik bebas
                strict = f["ref"] not in ("authweb", "whsasg", "pccatasg")
                dv = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=strict,
                                    errorStyle="stop" if strict else "warning")
                dv.add(f"{letter}3:{letter}{ROWS + 2}")
                ws.add_data_validation(dv)
        for r, row in enumerate(samples, 3):
            for c, v in enumerate(row, 1):
                cell = ws.cell(r, c, "" if v is None else str(v))
                cell.border = BORDER
        ws.freeze_panes = "B3"
        ws.row_dimensions[1].height = 30

    # sheet Keterangan
    ws = wb.create_sheet("Keterangan")
    r = 1
    ws.cell(r, 1, title).font = Font(bold=True, size=14, color="1F4E78"); r += 1
    ws.cell(r, 1, f"Sumber struktur: {refs['company']} ({SCHEMA} @ {HOST}). Dibuat {date.today():%Y-%m-%d}.").font = NOTE_FONT; r += 2
    ws.cell(r, 1, "Fungsi").font = Font(bold=True, color="1F4E78"); r += 1
    for line in purpose:
        ws.cell(r, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 15 * (1 + len(line) // 140)
        r += 1
    r += 1
    for c, h in enumerate(["Sheet", "Kolom", "Field", "Wajib", "Cara isi"], 1):
        cell = ws.cell(r, c, h)
        cell.fill, cell.font, cell.border = LABEL_FILL, LABEL_FONT, BORDER
    r += 1
    for name, fields, _ in sheets:
        for f in fields:
            for c, v in enumerate([name, f["label"], f["field"], "Ya" if f["req"] else "", f["note"]], 1):
                cell = ws.cell(r, c, v)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    r += 1
    for line in ["Baris 1 = label (tanda * = wajib), baris 2 = nama field (jangan diubah), isi data mulai baris 3. Hapus baris contoh sebelum upload.",
                 "Kolom berdropdown: pilih dari daftar (lihat sheet Referensi). Daftar diambil dari SAP saat template dibuat.", *extra_notes]:
        ws.cell(r, 1, line).font = NOTE_FONT
        r += 1
    for c, w in enumerate([16, 30, 28, 8, 95], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # isi sheet Referensi
    ws = wb.create_sheet("Referensi")
    for col, head, cols, rows in ref_blocks:
        ws.cell(1, col, head).font = Font(bold=True, color="1F4E78")
        for i, h in enumerate(cols):
            cell = ws.cell(2, col + i, h)
            cell.fill, cell.font, cell.border = FIELD_FILL, FIELD_FONT, BORDER
            ws.column_dimensions[get_column_letter(col + i)].width = 14 if i == 0 else 32
        for r, row in enumerate(rows, 3):
            for i, v in enumerate(row[:len(cols)]):
                ws.cell(r, col + i, str(v)).border = BORDER
        ws.column_dimensions[get_column_letter(col + len(cols))].width = 3
    ws.freeze_panes = "A3"

    wb.save(path)
    print("OK", path)


def main():
    refs = fetch_refs()
    prefix = f"iReapBiz {SCHEMA}"
    out = HERE

    # ---------- 01 Petty Cash Category Assignment ----------
    # contoh: kategori transport yang sudah ada di DB, dikumpulkan dalam satu dokumen "Transport"
    pc_lines = [("1", i, a, n) for i, (_, _, a, n) in enumerate(refs["pccat_lines"], 1)] or [("1", 1, "6101001", "Transport-BBM")]
    build(out / f"{prefix} - 01 - Petty Cash Category Assignment.xlsx", "Template Petty Cash Category Assignment (UDO STEM_PCCATASG)", [
        "Daftar kategori biaya kas kecil yang boleh dipilih seorang user di web iReapBiz saat membuat Petty Cash Request dan Cash Bon Claim.",
        "Setiap kategori = satu akun G/L biaya (contoh Transport-BBM = 6101001 BIAYA BAHAN BAKAR). Saat cash bon diselesaikan, jurnal settlement men-debit akun kategori dan meng-kredit akun kas kecil (PettyCash Acc Web di Employee Master Data).",
        "Satu dokumen assignment bisa dipakai banyak karyawan: dokumen dihubungkan ke karyawan lewat Petty Cash Cat Asg No. (DocEntry) dan Document Number (DocNum) di tab iReapBiz Employee Master Data. Contoh: dokumen 'Transport' untuk kurir, dokumen 'Operasional Kantor' untuk admin.",
    ], [
        ("Header", [
            F("No. Dokumen", "DocNum", "Nomor urut dokumen di template ini (1, 2, 3 ...), hanya untuk menghubungkan Header dengan Detail. Nomor final dibuat SAP saat import.", req=True),
            F("Keterangan", "U_Remarks", "Nama / keterangan assignment, mis. 'Transport', 'Operasional Kantor'. Maks 254 karakter.", req=True, width=40),
        ], [("1", "Transport")]),
        ("Detail", [
            F("No. Dokumen", "DocNum", "Sama dengan No. Dokumen di sheet Header.", req=True),
            F("No. Baris", "LineId", "Nomor baris per dokumen, mulai 1 dan berurutan.", req=True),
            F("Akun G/L", "U_AcctCode", "Kode akun biaya (posting account) untuk kategori ini. Pilih dari daftar. Nilai inilah yang tersimpan sebagai kategori di Petty Cash Request.", req=True, ref="accounts"),
            F("Nama Kategori", "U_PettyCashCatName", "Nama kategori yang tampil di web, mis. 'Transport-BBM'. Maks 100 karakter.", req=True, width=32),
        ], pc_lines),
    ], refs, {"accounts": ("Akun G/L (posting)", ["AcctCode", "AcctName"]),
              "pccatasg": ("Dokumen yang sudah ada", ["DocEntry", "DocNum", "U_Remarks"])},
        ["Satu akun G/L sebaiknya tidak dipakai dua kali dalam satu dokumen."])

    # ---------- 02 Warehouse Assignment ----------
    # contoh: dokumen 1 = semua gudang, dokumen 2 = satu gudang saja
    whs = refs["warehouses"] or [("HO", "Head Office")]
    wl = [("1", i, c, n) for i, (c, n) in enumerate(whs, 1)] + [("2", 1, *whs[0])]
    build(out / f"{prefix} - 02 - Warehouse Assignment.xlsx", "Template Warehouse Assignment (UDO STEM_WHSASGNWEBP)", [
        "Daftar gudang yang boleh diakses seorang user di web iReapBiz. Pilihan gudang di transaksi web (mis. Sales Quotation / Order, Delivery, transfer, stok) hanya menampilkan gudang yang ada di assignment user tersebut.",
        "Satu dokumen assignment = satu profil akses gudang (contoh 'All warehouse' berisi semua gudang, 'BAL-1' hanya gudang Bali 1) dan bisa dipakai banyak karyawan.",
        "Dokumen dihubungkan ke karyawan lewat Warehouse Assignment No. (DocEntry) dan Warehouse Assignment Document Number (DocNum) di tab iReapBiz Employee Master Data.",
    ], [
        ("Header", [
            F("No. Dokumen", "DocNum", "Nomor urut dokumen di template ini (1, 2, 3 ...), hanya untuk menghubungkan Header dengan Detail. Nomor final dibuat SAP saat import.", req=True),
            F("Keterangan", "U_Remarks", "Nama profil akses, mis. 'All warehouse', 'Gudang Bali'. Maks 254 karakter.", req=True, width=40),
            F("Branch", "U_BRANCH", "Kode branch / area (teks bebas, maks 20 karakter). Database ini tidak memakai multiple branch, boleh dikosongkan."),
        ], [("1", "All warehouse", ""), ("2", f"Gudang {whs[0][1]}", "")]),
        ("Detail", [
            F("No. Dokumen", "DocNum", "Sama dengan No. Dokumen di sheet Header.", req=True),
            F("No. Baris", "LineId", "Nomor baris per dokumen, mulai 1 dan berurutan.", req=True),
            F("Kode Gudang", "U_WhsCode", "Kode gudang yang boleh diakses. Pilih dari daftar.", req=True, ref="warehouses"),
            F("Nama Gudang", "U_WhsName", "Nama gudang, samakan dengan nama di SAP (lihat sheet Referensi). Maks 100 karakter.", req=True, width=28),
        ], wl),
    ], refs, {"warehouses": ("Gudang aktif", ["WhsCode", "WhsName"]),
              "whsasg": ("Dokumen yang sudah ada", ["DocEntry", "DocNum", "U_Remarks"])})

    # ---------- 03 Employee Master Data ----------
    dist = {f"dim{d}": (f"Dist. Rule {d}", ["OcrCode", "OcrName"]) for d in range(1, 6)}
    build(out / f"{prefix} - 03 - Employee Master Data.xlsx", "Template Employee Master Data (OHEM) - tab iReapBiz", [
        "Data karyawan sekaligus akun login web iReapBiz. Kolom U_ mengikuti tab iReapBiz di Employee Master Data.",
        "Isi Warehouse Assignment dan Petty Cash Category Assignment lebih dulu (template 02 dan 01), lalu isi nomor dokumennya di sini.",
    ], [("Employee", [
        F("Employee Code", "EmployeeID", "Nomor karyawan (empID). Untuk karyawan baru kosongkan, SAP membuat nomornya otomatis. Isi hanya untuk update karyawan yang sudah ada."),
        F("First Name", "FirstName", "Nama depan.", req=True),
        F("Last Name", "LastName", "Nama belakang. Isi '-' bila tidak ada.", req=True),
        F("User Web", "U_ST_WEB_USER", "Username login web iReapBiz, harus unik. Maks 100 karakter.", req=True),
        F("Password Web", "U_ST_WEB_PASS", "Kosongkan: SAP mengisi password default (hash). Ganti lewat tombol Reset Password di tab iReapBiz. Jangan isi password asli di file ini."),
        F("SAP Indirect User Web", "U_ST_SAPB1_USER", "User SAP B1 (indirect access) yang dipakai web untuk posting dokumen. Maks 25 karakter."),
        F("SAP Indirect Password Web", "U_ST_SAPB1_PASS", "Password user SAP di atas. File ini jadi berisi password: jangan disebar, hapus setelah import."),
        F("Authorization Code", "U_STEM_IBAuthID", "DocEntry dokumen Authorization WEB (hak akses menu web). Pilih dari daftar.", req=True, ref="authweb"),
        F("Authorization Code Number", "U_STEM_IBAuthIDNum", "DocNum dokumen Authorization WEB yang sama (lihat sheet Referensi).", req=True),
        F("Warehouse Assignment No.", "U_ST_WhsAssignmentNo", "DocEntry dokumen Warehouse Assignment (template 02). Pilih dari daftar, atau ketik nomor dokumen yang baru di-import.", req=True, ref="whsasg"),
        F("Warehouse Assignment Document Number", "U_ST_WhsAssignmentDocNum", "DocNum dokumen Warehouse Assignment yang sama."),
        F("Petty Cash Cat Asg No.", "U_ST_PCCatCode", "DocEntry dokumen Petty Cash Category Assignment (template 01). Pilih dari daftar, atau ketik nomor dokumen yang baru di-import.", ref="pccatasg"),
        F("Petty Cash Cat Asg Document Number", "U_ST_PCCatNum", "DocNum dokumen Petty Cash Category Assignment yang sama."),
        F("Is Reseller", "U_STEM_ISRESELLER", "Y = karyawan ini reseller, N = bukan. Kosong = N.", enum=["Y", "N"]),
        *[F(f"Dist. Rule {d} WEB", f"U_ST_DistRule{d}", f"Distribution rule (dimension {d}) default untuk transaksi web user ini. Pilih dari daftar.", ref=f"dim{d}") for d in range(1, 6)],
        F("PettyCash Acc Web", "U_ST_PCAccount", "Akun kas kecil (mis. 1101002 KAS KECIL) yang dikredit saat cash bon / petty cash user ini. Pilih dari daftar.", ref="accounts"),
    ], [("", "BUDI", "SANTOSO", "BUDI01", "", "", "", "2", "1", "2", "1", "3", "1", "N", "BAL", "", "", "", "", "1101002")])],
        refs, {"authweb": ("Authorization WEB", ["DocEntry", "DocNum", "U_Remarks"]),
               "whsasg": ("Warehouse Assignment", ["DocEntry", "DocNum", "U_Remarks"]),
               "pccatasg": ("Petty Cash Cat Asg", ["DocEntry", "DocNum", "U_Remarks"]),
               **dist, "accounts": ("Akun G/L (posting)", ["AcctCode", "AcctName"])},
        ["No. (DocEntry) dan Document Number (DocNum) dokumen UDO biasanya sama, tapi bisa berbeda bila ada dokumen yang dihapus atau series berbeda: cek di sheet Referensi."])


if __name__ == "__main__":
    main()
